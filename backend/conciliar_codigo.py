#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
conciliar_codigo.py — Conciliación por CÓDIGO AIR26 + importe (el método exacto).

Para cuando el export del banco trae el CONCEPTO real (columna con el código
AIR26-XXXXX que el cliente escribió), no solo el nombre del ordenante. Es el
método bueno: casa por código como el backend (`conciliarBanco` en Code.gs) y no
depende del cruce por nombre, así que desaparecen los ambiguos.

Hermano de casador.py:
  · casador.py         → cuando NO hay código, casa por NOMBRE + importe.
  · conciliar_codigo.py→ cuando SÍ hay código, casa por CÓDIGO + importe (esto).

Para las pocas transferencias SIN código detectable (donativos con mensaje,
"AIR00785" mal escrito, Revolut…) hace fallback por nombre reutilizando el casador.

Entrada (--banco): xlsx del banco con columnas reconocidas por su cabecera:
  ORDENANTE (nombre) · CONCEPTO (código) · IMPORTE · (FECHA opcional).
--pedidos: export de la hoja PEDIDOS (.xlsx con hoja PEDIDOS, o .csv).

Uso:
  python conciliar_codigo.py --banco transferencias.xlsx \
                             --pedidos "Plataforma Solidaria ....xlsx" \
                             --salida  conciliacion-resultado.xlsx

Cubos del RESULTADO:
  CASADO           código ok, pedido no pagado, importe exacto → confirmar
  CASADO_NOMBRE    sin código pero casado por nombre+importe (fallback)
  YA_PAGADO        código ok, el pedido ya estaba pagado
  REVISAR_IMPORTE  código ok pero el importe no coincide con el total del pedido
  REVISAR_CODIGO_INEXISTENTE  código detectado que no existe en PEDIDOS
  REVISAR_CADUCADO / REVISAR_ANULADO  el pedido está fuera de juego (ver CONCILIACION.md §7b)
  REVISAR_AMBIGUO  sin código; varios pedidos posibles por nombre
  SIN_MATCH        sin código y sin pedido por nombre → donativo o cuenta de tercero
"""
import argparse
import re
import sys
from difflib import SequenceMatcher

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import casador as C   # tokens_clave, compacto, parse_importe (mismo directorio)

# Espejo de Code.gs: estados pagados y "fuera de juego".
PAGADOS = {"PAGO_CONCILIADO", "ENVIADO_PROVEEDOR", "RECIBIDO", "LISTO_RECOGIDA", "ENTREGADO"}
FUERA = {"CADUCADO", "ANULADO"}
UMBRAL_FUZZY = 0.86

COLORES = {
    "CASADO": "C6EFCE", "CASADO_NOMBRE": "C6EFCE", "YA_PAGADO": "DDEBF7",
    "REVISAR_IMPORTE": "FFEB9C", "REVISAR_AMBIGUO": "FFEB9C",
    "REVISAR_CADUCADO": "F4B183", "REVISAR_ANULADO": "D9A0A0",
    "REVISAR_CODIGO_INEXISTENTE": "FFC7CE", "SIN_MATCH": "FFC7CE",
}
ORDEN = {"CASADO": 0, "CASADO_NOMBRE": 1, "REVISAR_IMPORTE": 2,
         "REVISAR_CODIGO_INEXISTENTE": 3, "REVISAR_CADUCADO": 4, "REVISAR_ANULADO": 5,
         "REVISAR_AMBIGUO": 6, "SIN_MATCH": 7, "YA_PAGADO": 8}


def detectar_id(concepto, prefijo="AIR26"):
    """Espejo de detectarId() del Code.gs: aplana y busca PREFIJO + número → AIR26-XXXXX.
    Tolera 'AIR26-00123', 'AIR26 00123', 'AIR2600123', etc. Devuelve '' si no hay."""
    pref = re.sub(r"[^A-Z0-9]", "", str(prefijo).upper())
    flat = re.sub(r"[^A-Z0-9]", "", str(concepto).upper())
    m = re.search(pref + r"0*(\d{1,6})", flat)
    return "AIR26-" + str(int(m.group(1))).zfill(5) if m else ""


def leer_banco(path):
    """xlsx del banco. Reconoce ORDENANTE / CONCEPTO / IMPORTE / FECHA por cabecera."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c or "").strip().upper() for c in filas[0]]

    def idx(*claves):
        for k in claves:
            for i, c in enumerate(cab):
                if k in c:
                    return i
        return None

    i_ord = idx("ORDENANTE", "NOMBRE", "TITULAR")
    i_con = idx("CONCEPTO")
    i_imp = idx("IMPORTE")
    i_fec = idx("FECHA")
    if i_con is None or i_imp is None:
        sys.exit("El export del banco no tiene columnas CONCEPTO / IMPORTE reconocibles.\n"
                 "Cabecera leída: " + ", ".join(cab))
    out = []
    for r in filas[1:]:
        if r is None:
            continue
        con = r[i_con] if i_con < len(r) else None
        ordn = r[i_ord] if i_ord is not None and i_ord < len(r) else ""
        if (con in (None, "")) and (ordn in (None, "")):
            continue
        out.append({
            "ordenante": str(ordn or ""),
            "concepto": str(con or ""),
            "importe": C.parse_importe(r[i_imp] if i_imp < len(r) else 0),
            "fecha": r[i_fec] if i_fec is not None and i_fec < len(r) else "",
        })
    return out


def leer_pedidos(path):
    """Reutiliza el lector del casador y añade índices por ID y por importe."""
    peds = C.leer_pedidos(path)
    por_id, por_importe = {}, {}
    for p in peds:
        por_id[p["id"]] = p
        por_importe.setdefault(C.cent(p["total"]), []).append(p)
    return peds, por_id, por_importe


def conciliar(banco, por_id, por_importe):
    res = []
    for t in banco:
        con, imp = t["concepto"], t["importe"]
        code = detectar_id(con)
        if code:
            p = por_id.get(code)
            if not p:
                res.append(_fila(t, "REVISAR_CODIGO_INEXISTENTE", nota="código " + code + " no existe en PEDIDOS"))
            elif p["estado"] in PAGADOS:
                res.append(_fila(t, "YA_PAGADO", p))
            elif p["estado"] in FUERA:
                res.append(_fila(t, "REVISAR_" + p["estado"], p,
                                 nota="pedido " + p["estado"] + " — reactivar o donativo (§7b)"))
            elif C.cent(imp) == C.cent(p["total"]):
                res.append(_fila(t, "CASADO", p))
            else:
                res.append(_fila(t, "REVISAR_IMPORTE", p,
                                 nota="banco %.2f vs pedido %.2f" % (imp, p["total"])))
            continue
        # --- sin código: fallback por NOMBRE + importe (casador) ---
        res.append(_fallback_nombre(t, por_importe))
    return res


def _fallback_nombre(t, por_importe):
    btok = C.tokens_clave(t["ordenante"])
    bcomp = C.compacto(t["ordenante"])
    cand = por_importe.get(C.cent(t["importe"]), [])
    ex = [p for p in cand if p["tokens"] == btok and btok]
    if not ex:
        ex = [p for p in cand if btok and (btok <= p["tokens"] or p["tokens"] <= btok)
              and min(len(btok), len(p["tokens"])) >= 2]
    if len(ex) == 1:
        p = ex[0]
        if p["estado"] in PAGADOS:
            return _fila(t, "YA_PAGADO", p, nota="sin código; casado por nombre")
        if p["estado"] in FUERA:
            return _fila(t, "REVISAR_" + p["estado"], p, nota="sin código; " + p["estado"])
        return _fila(t, "CASADO_NOMBRE", p, nota="sin código; casado por nombre+importe")
    if len(ex) > 1:
        return _fila(t, "REVISAR_AMBIGUO", ex[0], nota="sin código; %d pedidos mismo nombre/importe" % len(ex))
    aprox = sorted(((SequenceMatcher(None, bcomp, p["compacto"]).ratio(), p) for p in cand),
                   key=lambda x: -x[0])
    if aprox and aprox[0][0] >= UMBRAL_FUZZY:
        return _fila(t, "REVISAR_AMBIGUO", aprox[0][1], nota="sin código; parecido %.0f%%" % (aprox[0][0] * 100))
    return _fila(t, "SIN_MATCH", nota="sin código y sin pedido por nombre — donativo o cuenta de tercero")


def _fila(t, resultado, p=None, nota=""):
    return {
        "ordenante": t["ordenante"], "fecha": t["fecha"], "concepto": t["concepto"],
        "importe": t["importe"], "resultado": resultado,
        "pid": p["id"] if p else "", "pnom": p["nombre_completo"] if p else "",
        "pemail": p["email"] if p else "", "ptot": p["total"] if p else "", "nota": nota,
    }


def escribir(res, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliación"
    cols = ["ORDENANTE", "FECHA", "CONCEPTO", "IMPORTE", "RESULTADO",
            "PEDIDO_ID", "NOMBRE_PEDIDO", "EMAIL", "TOTAL_PEDIDO", "NOTA"]
    ws.append(cols)
    hfill = PatternFill("solid", fgColor="1F3A5F")
    for i in range(1, len(cols) + 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
    res = sorted(res, key=lambda r: ORDEN.get(r["resultado"], 9))
    for r in res:
        ws.append([r["ordenante"], r["fecha"], r["concepto"], r["importe"], r["resultado"],
                   r["pid"], r["pnom"], r["pemail"], r["ptot"], r["nota"]])
        row = ws.max_row
        if r["resultado"] in COLORES:
            ws.cell(row, 5).fill = PatternFill("solid", fgColor=COLORES[r["resultado"]])
        ws.cell(row, 4).number_format = "#,##0.00 €"
        if r["ptot"] != "":
            ws.cell(row, 9).number_format = "#,##0.00 €"
        ws.cell(row, 2).number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    for i, w in enumerate([30, 12, 26, 12, 26, 14, 30, 28, 13, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="Conciliación banco↔pedidos por CÓDIGO AIR26 + importe")
    ap.add_argument("--banco", required=True, help="xlsx del banco (ORDENANTE, CONCEPTO, IMPORTE)")
    ap.add_argument("--pedidos", required=True, help="Export de la hoja PEDIDOS (.xlsx o .csv)")
    ap.add_argument("--salida", default="conciliacion-resultado.xlsx")
    a = ap.parse_args()

    banco = leer_banco(a.banco)
    _, por_id, por_importe = leer_pedidos(a.pedidos)
    res = conciliar(banco, por_id, por_importe)
    escribir(res, a.salida)

    from collections import Counter
    cnt = Counter(r["resultado"] for r in res)
    suma = Counter()
    for r in res:
        suma[r["resultado"]] += r["importe"]
    print(f"Transferencias : {len(banco)}")
    print(f"Suma importes  : {round(sum(r['importe'] for r in res), 2)} €")
    print("-" * 48)
    for k in sorted(cnt, key=lambda x: ORDEN.get(x, 9)):
        print(f"  {k:<28} {cnt[k]:>4}   {round(suma[k], 2):>10.2f} €")
    print("-" * 48)
    print(f"Salida → {a.salida}")


if __name__ == "__main__":
    main()
